import re
from typing import Dict, Any
import pymupdf as fitz
import docx
import io


def extract_contact_info(text: str) -> Dict[str, Any]:
    email_match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text)
    phone_match = re.search(r"(?:\+\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)?\d{3,4}[-.\s]?\d{4}|(?:\+\d{1,3}[-.\s]?)?\d{3}[-.\s]\d{4}", text)
    linkedin_match = re.search(r"linkedin\.com/in/[\w\-]+", text, re.IGNORECASE)
    github_match = re.search(r"github\.com/[\w\-]+", text, re.IGNORECASE)

    return {
        "email": email_match.group(0) if email_match else "",
        "phone": phone_match.group(0) if phone_match else "",
        "linkedin": linkedin_match.group(0) if linkedin_match else "",
        "github": github_match.group(0) if github_match else ""
    }


def parse_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Extracts text and evaluates document structure.
    Raises ValueError on empty or unparseable files without masking errors.
    """
    if not file_bytes:
        raise ValueError("Provided file is empty (0 bytes).")

    fn_lower = filename.lower()
    raw_text = ""
    page_count = 1
    is_multi_column = False
    has_tables = False

    if fn_lower.endswith(".pdf"):
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            page_count = len(doc)
            pages_text = []

            for page in doc:
                text = page.get_text()
                pages_text.append(text)

                blocks = page.get_text("blocks")
                left_cols, right_cols = 0, 0
                width = page.rect.width
                for b in blocks:
                    if b[0] < width * 0.45:
                        left_cols += 1
                    elif b[0] > width * 0.50:
                        right_cols += 1

                if left_cols >= 3 and right_cols >= 3:
                    is_multi_column = True

                tables = page.find_tables()
                if tables.tables:
                    has_tables = True

            raw_text = "\n".join(pages_text).strip()
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

    elif fn_lower.endswith((".docx", ".doc")):
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            has_tables = len(doc.tables) > 0
            raw_text = "\n".join(paragraphs).strip()
            page_count = max(1, len(raw_text) // 2500)
        except Exception as e:
            raise ValueError(f"Failed to parse Word document: {str(e)}")

    else:
        try:
            raw_text = file_bytes.decode("utf-8", errors="replace").strip()
            page_count = max(1, len(raw_text) // 2500)
        except Exception as e:
            raise ValueError(f"Failed to parse plain text file: {str(e)}")

    if not raw_text:
        raise ValueError("Could not extract any readable text from the document. Ensure it is not a scanned raster image without OCR.")

    contact = extract_contact_info(raw_text)

    # Basic section detection
    sections = {}
    section_headers = ["experience", "education", "skills", "projects", "certifications", "summary"]
    for header in section_headers:
        match = re.search(rf"\b{header}\b", raw_text, re.IGNORECASE)
        sections[header] = match is not None

    return {
        "raw_text": raw_text,
        "page_count": page_count,
        "is_multi_column": is_multi_column,
        "has_tables": has_tables,
        "contact": contact,
        "sections_detected": sections,
        "char_count": len(raw_text),
        "word_count": len(raw_text.split())
    }
